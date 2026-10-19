"""Excel tools: sort workbooks, and pull tables out of Word/PDF/image files into .xlsx."""

import bisect
import copy
import datetime
import io
import re
import statistics

import cv2  # opencv-python-headless, installed with pdf2docx
import docx
import fitz  # PyMuPDF
import numpy as np
import pytesseract
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image, ImageOps

from pdf_tools import _configure_tesseract

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E8E9FB")
_MAX_COLUMN_WIDTH = 60
_MAX_SHEET_NAME = 31  # Excel's hard limit

# "1,234.50", "-12", "(12)", "$1,000", "R 250", "€3.5", "45%" - but not IDs like "00123".
_NUMBER_RE = re.compile(r"^\(?-?[$€£R]?\s?-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\)?%?$")


# ---------------------------------------------------------------------------
# Writing extracted tables
# ---------------------------------------------------------------------------

def _to_number(text: str):
    """Turn a numeric-looking string into a number so the sheet can be sorted
    and summed. Returns (value, number_format) - or the original text if it
    isn't clearly a number."""
    s = text.strip()
    if not s or not _NUMBER_RE.match(s):
        return text, None
    digits = re.sub(r"[^\d.]", "", s)
    if len(digits) > 1 and digits[0] == "0" and not digits.startswith("0."):
        return text, None  # leading zero: an ID/code, keep it as text
    if len(digits.replace(".", "")) > 15:
        return text, None  # beyond Excel's precision (account/card numbers)
    try:
        value = float(digits)
    except ValueError:
        return text, None
    if s.startswith("(") and s.endswith(")") or "-" in s:
        value = -value
    if s.endswith("%"):
        return value / 100, "0.##%"
    if value.is_integer() and "." not in digits:
        value = int(value)
    if "," in s:
        return value, "#,##0.00" if "." in digits else "#,##0"
    return value, None


def _clean_cell(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _unique_sheet_name(name: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "-", name)[:_MAX_SHEET_NAME] or "Sheet"
    candidate, counter = name, 2
    while candidate.casefold() in used:
        suffix = f" ({counter})"
        candidate = name[: _MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    used.add(candidate.casefold())
    return candidate


def write_tables_xlsx(tables: list[tuple[str, list[list[str]]]], out_path: str) -> None:
    """Write each (sheet name, rows) table to its own worksheet."""
    tables = [(name, rows) for name, rows in tables if any(any(c for c in r) for r in rows)]
    if not tables:
        raise ValueError("No tables were found in that file.")

    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    for name, rows in tables:
        ws = wb.create_sheet(_unique_sheet_name(name, used_names))
        widths: dict[int, int] = {}
        for r, row in enumerate(rows, start=1):
            for c, raw in enumerate(row, start=1):
                text = _clean_cell(raw)
                if not text:
                    continue
                # The header row stays text even when it looks numeric (e.g. years).
                value, number_format = (text, None) if r == 1 else _to_number(text)
                cell = ws.cell(row=r, column=c, value=value)
                if number_format:
                    cell.number_format = number_format
                widths[c] = max(widths.get(c, 0), len(text))

        n_cols = max((len(row) for row in rows), default=0)
        if len(rows) > 1:
            for c in range(1, n_cols + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
        for c in range(1, n_cols + 1):
            ws.column_dimensions[get_column_letter(c)].width = min(
                max(widths.get(c, 0) + 2, 8), _MAX_COLUMN_WIDTH
            )
    wb.save(out_path)


# ---------------------------------------------------------------------------
# Word -> Excel
# ---------------------------------------------------------------------------

def _docx_table_rows(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        seen = set()
        cells = []
        for cell in row.cells:
            # python-docx repeats a horizontally merged cell once per grid
            # column it spans - keep its text in the first column only.
            if id(cell._tc) in seen:
                cells.append("")
                continue
            seen.add(id(cell._tc))
            cells.append(cell.text)
        rows.append(cells)
    return rows


def _iter_docx_tables(tables):
    for table in tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_docx_tables(cell.tables)


def docx_tables_to_xlsx(src: str, out_path: str) -> None:
    document = docx.Document(src)
    tables = [
        (f"Table {i}", _docx_table_rows(t))
        for i, t in enumerate(_iter_docx_tables(document.tables), start=1)
    ]
    write_tables_xlsx(tables, out_path)


# ---------------------------------------------------------------------------
# Image -> Excel (OCR)
# ---------------------------------------------------------------------------

def _prepare_image(img: Image.Image) -> Image.Image:
    img = ImageOps.exif_transpose(img)
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        background = Image.new("RGBA", img.size, "white")
        img = Image.alpha_composite(background, img)
    img = img.convert("L")
    if img.width < 1500:
        img = img.resize((img.width * 2, img.height * 2), Image.LANCZOS)
    return ImageOps.autocontrast(img)


def _line_positions(mask: np.ndarray, axis: int, min_length: float) -> list[int]:
    """Centres of the ruling lines in a line mask (axis=1: horizontal lines)
    that are at least min_length pixels long."""
    profile = mask.sum(axis=axis) / 255
    if not profile.any():
        return []
    hits = np.flatnonzero(profile >= max(profile.max() * 0.5, min_length))
    if not len(hits):
        return []
    groups = np.split(hits, np.flatnonzero(np.diff(hits) > 3) + 1)
    return [int(g.mean()) for g in groups]


def _remove_grid(img: Image.Image) -> tuple[Image.Image, list[int], list[int]]:
    """Find table ruling lines, erase them (Tesseract reads nothing through a
    dense grid) and return where they were so they can mark cell edges.

    Only long rules count - tall letters like "d" or "l" would otherwise pass
    as vertical lines and get erased. Without at least two horizontal rules
    the image is left untouched."""
    gray = np.array(img)
    _, binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    h, w = binary.shape
    horizontal = cv2.morphologyEx(
        binary, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(80, w // 10), 1))
    )
    ys = _line_positions(horizontal, 1, min_length=max(150, w * 0.2))
    if len(ys) < 2:
        return img, [], []

    # A real column rule spans most of the table between the outer rows.
    span = ys[-1] - ys[0]
    vertical = cv2.morphologyEx(
        binary, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(80, int(span * 0.4))))
    )
    xs = _line_positions(vertical, 0, min_length=span * 0.5)
    if len(xs) < 2:
        xs, vertical = [], np.zeros_like(vertical)

    lines = cv2.dilate(horizontal | vertical, np.ones((3, 3), np.uint8))
    gray[lines > 0] = 255
    return Image.fromarray(gray), ys, xs


def _grid_table(words: list[dict], ys: list[int], xs: list[int]) -> list[list[str]]:
    """Drop each word into the grid cell its centre falls in."""
    cells: dict[tuple[int, int], list[dict]] = {}
    for word in words:
        cx = (word["x0"] + word["x1"]) / 2
        r = bisect.bisect(ys, word["cy"]) - 1
        c = bisect.bisect(xs, cx) - 1
        if 0 <= r < len(ys) - 1 and 0 <= c < len(xs) - 1:
            cells.setdefault((r, c), []).append(word)
    table = []
    for r in range(len(ys) - 1):
        row = []
        for c in range(len(xs) - 1):
            in_cell = sorted(cells.get((r, c), []), key=lambda w: (round(w["cy"] / max(w["h"], 1)), w["x0"]))
            row.append(" ".join(w["text"] for w in in_cell))
        if any(row):
            table.append(row)
    return table


def _ocr_words(img: Image.Image) -> list[dict]:
    data = pytesseract.image_to_data(img, config="--psm 6", output_type=pytesseract.Output.DICT)
    words = []
    for i, text in enumerate(data["text"]):
        text = text.strip()
        try:
            conf = float(data["conf"][i])
        except (TypeError, ValueError):
            conf = -1
        if not text or conf < 30:
            continue
        x, y, w, h = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
        words.append({"text": text, "x0": x, "x1": x + w, "cy": y + h / 2, "h": h})
    return words


def _group_rows(words: list[dict], median_h: float) -> list[list[dict]]:
    rows: list[list[dict]] = []
    for word in sorted(words, key=lambda w: w["cy"]):
        if rows:
            row = rows[-1]
            row_cy = sum(w["cy"] for w in row) / len(row)
            if abs(word["cy"] - row_cy) <= median_h * 0.6:
                row.append(word)
                continue
        rows.append([word])
    return [sorted(row, key=lambda w: w["x0"]) for row in rows]


def _split_cells(row: list[dict], gap: float) -> list[dict]:
    """Join neighbouring words into cells; a wide horizontal gap starts a new cell."""
    cells = []
    for word in row:
        if cells and word["x0"] - cells[-1]["x1"] < gap:
            cells[-1]["text"] += " " + word["text"]
            cells[-1]["x1"] = max(cells[-1]["x1"], word["x1"])
        else:
            cells.append({"text": word["text"], "x0": word["x0"], "x1": word["x1"]})
    return cells


def _column_bands(rows_of_cells: list[list[dict]]) -> list[tuple[int, int]]:
    """Find column bands from where cells line up across rows. Single-cell
    rows (titles, notes) are ignored, and an x-position must be covered by a
    few rows before it counts - so one stray wide cell can't fuse columns."""
    multi = [cells for cells in rows_of_cells if len(cells) > 1] or rows_of_cells
    width = max(c["x1"] for cells in multi for c in cells) + 1
    coverage = [0] * width
    for cells in multi:
        for c in cells:
            for x in range(c["x0"], c["x1"]):
                coverage[x] += 1
    threshold = max(1, round(len(multi) * 0.15))
    bands, start = [], None
    for x, count in enumerate(coverage + [0]):
        if count >= threshold and start is None:
            start = x
        elif count < threshold and start is not None:
            bands.append((start, x))
            start = None
    return bands


def _best_band(cell: dict, bands: list[tuple[int, int]]) -> int:
    def overlap(band):
        return min(cell["x1"], band[1]) - max(cell["x0"], band[0])

    best = max(range(len(bands)), key=lambda i: overlap(bands[i]))
    if overlap(bands[best]) > 0:
        return best
    centre = (cell["x0"] + cell["x1"]) / 2
    return min(range(len(bands)), key=lambda i: abs((bands[i][0] + bands[i][1]) / 2 - centre))


def image_to_table(img: Image.Image) -> list[list[str]]:
    """OCR an image and rebuild its rows/columns from where the words sit."""
    _configure_tesseract()
    cleaned, ys, xs = _remove_grid(_prepare_image(img))
    words = _ocr_words(cleaned)
    if not words:
        return []
    if len(ys) >= 2 and len(xs) >= 2:
        table = _grid_table(words, ys, xs)
        if table:
            return table

    median_h = statistics.median(w["h"] for w in words)
    rows_of_cells = [_split_cells(row, gap=median_h * 1.2) for row in _group_rows(words, median_h)]
    bands = _column_bands(rows_of_cells)
    if not bands:
        return [[c["text"] for c in cells] for cells in rows_of_cells]

    table = []
    for cells in rows_of_cells:
        out = [""] * len(bands)
        for cell in cells:
            i = _best_band(cell, bands)
            out[i] = f"{out[i]} {cell['text']}".strip()
        table.append(out)
    return table


def image_table_to_xlsx(src: str, out_path: str) -> None:
    with Image.open(src) as img:
        rows = image_to_table(img)
    write_tables_xlsx([("Table 1", rows)], out_path)


# ---------------------------------------------------------------------------
# PDF -> Excel
# ---------------------------------------------------------------------------

def _looks_like_real_table(rows: list[list]) -> bool:
    """The text strategy finds 'tables' in ordinary paragraphs too - only keep
    results that have at least 2x2 cells and are mostly filled in."""
    if len(rows) < 2 or max(len(r) for r in rows) < 2:
        return False
    cells = [c for r in rows for c in r]
    filled = sum(1 for c in cells if c not in (None, ""))
    return filled >= len(cells) * 0.5


def pdf_tables_to_xlsx(src: str, out_path: str) -> None:
    tables = []
    with fitz.open(src) as doc:
        for page_no, page in enumerate(doc, start=1):
            if not page.get_text().strip():
                # Scanned page: no text layer, so OCR a rendering of it.
                pix = page.get_pixmap(dpi=300)
                with Image.open(io.BytesIO(pix.tobytes("png"))) as img:
                    rows = image_to_table(img)
                if rows:
                    tables.append((f"P{page_no}-T1", rows))
                continue

            found = [t.extract() for t in page.find_tables(strategy="lines").tables]
            if not found:
                found = [
                    rows
                    for rows in (t.extract() for t in page.find_tables(strategy="text").tables)
                    if _looks_like_real_table(rows)
                ]
            for t_no, rows in enumerate(found, start=1):
                tables.append((f"P{page_no}-T{t_no}", rows))
    write_tables_xlsx(tables, out_path)


# ---------------------------------------------------------------------------
# Auto-sort a workbook
# ---------------------------------------------------------------------------

def _sort_key(value):
    """Numbers, then dates, then text, then booleans. Never compares mixed types."""
    if isinstance(value, bool):
        return (3, value)
    if isinstance(value, (int, float)):
        return (0, value)
    if isinstance(value, datetime.datetime):
        return (1, value)
    if isinstance(value, datetime.date):
        return (1, datetime.datetime.combine(value, datetime.time()))
    if isinstance(value, datetime.time):
        return (1, datetime.datetime.combine(datetime.date.min, value))
    return (2, str(value).casefold())


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _find_sort_column(ws, column: str, header_row: int | None, min_col: int, max_col: int) -> int | None:
    column = column.strip()
    if not column:
        return min_col
    if header_row is not None:
        for c in range(min_col, max_col + 1):
            header = ws.cell(row=header_row, column=c).value
            if header is not None and str(header).strip().casefold() == column.casefold():
                return c
    if re.fullmatch(r"[A-Za-z]{1,3}", column):
        return column_index_from_string(column.upper())
    return None


def _sort_sheet_rows(ws, ws_values, column: str, descending: bool, has_header: bool) -> str | None:
    """Sort a sheet's rows in place. Returns None once sorted, otherwise why
    not - an empty string when there was simply nothing to sort."""
    min_row, max_row = ws.min_row, ws.max_row
    min_col, max_col = ws.min_column, ws.max_column
    header_row = min_row if has_header else None
    first_data = min_row + 1 if has_header else min_row
    if first_data >= max_row:
        return ""

    for merged in ws.merged_cells.ranges:
        if merged.max_row >= first_data and merged.min_row <= max_row:
            return "it has merged cells in the data"

    sort_col = _find_sort_column(ws, column, header_row, min_col, max_col)
    if sort_col is None:
        return f'it has no "{column.strip()}" column'
    max_col = max(max_col, sort_col)

    rows = []
    for r in range(first_data, max_row + 1):
        cells = []
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cells.append((cell.value, cell._style, cell.hyperlink, cell.comment))
        # Formula cells sort by their last calculated result, when Excel saved one.
        key_value = ws_values.cell(row=r, column=sort_col).value
        if key_value is None:
            raw = ws.cell(row=r, column=sort_col).value
            key_value = None if isinstance(raw, str) and raw.startswith("=") else raw
        rows.append((r, key_value, cells, ws.row_dimensions[r].height))

    filled = [row for row in rows if not _is_blank(row[1])]
    blanks = [row for row in rows if _is_blank(row[1])]  # always kept at the bottom
    filled.sort(key=lambda row: _sort_key(row[1]), reverse=descending)

    for new_r, (old_r, _key, cells, height) in enumerate(filled + blanks, start=first_data):
        for offset, (value, style, hyperlink, comment) in enumerate(cells):
            col = min_col + offset
            target = ws.cell(row=new_r, column=col)
            if isinstance(value, str) and value.startswith("=") and new_r != old_r:
                # Shift relative references so =A5*B5 follows its row, like Excel does.
                origin = f"{get_column_letter(col)}{old_r}"
                value = Translator(value, origin=origin).translate_formula(
                    f"{get_column_letter(col)}{new_r}"
                )
            target.value = value
            target._style = style
            target.hyperlink = copy.copy(hyperlink) if hyperlink else None
            target.comment = copy.copy(comment) if comment else None
        ws.row_dimensions[new_r].height = height
    return None


def sort_workbook(
    src: str, out_path: str, mode: str = "rows", column: str = "",
    descending: bool = False, has_header: bool = True,
) -> list[str]:
    """Sort rows within every sheet and/or the sheet tabs. Returns notes about
    sheets that were left unsorted."""
    wb = load_workbook(src)
    notes = []

    if mode in ("rows", "both"):
        values_wb = load_workbook(src, data_only=True)
        sorted_any = False
        for ws in wb.worksheets:
            reason = _sort_sheet_rows(ws, values_wb[ws.title], column, descending, has_header)
            if reason:
                notes.append(f'Sheet "{ws.title}" was left unsorted because {reason}.')
            elif reason is None:
                sorted_any = True
        if not sorted_any and notes:
            raise ValueError(" ".join(notes))

    if mode in ("tabs", "both"):
        wb._sheets.sort(key=lambda s: s.title.casefold(), reverse=descending)
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = False
        wb.active = 0
        wb.active.sheet_view.tabSelected = True

    wb.save(out_path)
    return notes
